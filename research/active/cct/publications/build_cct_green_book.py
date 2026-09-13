from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "livre-vert-cct.md"
OUTPUT = HERE.parent / "output" / "docx" / "CCT-livre-vert.docx"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.extend([fldChar1, instrText, fldChar2])

def add_inline(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)

def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.25)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.35)
    sec.right_margin = Cm(2.15)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.07
    for name, size, before, after in (("Title", 25, 0, 14), ("Subtitle", 12.5, 0, 16), ("Heading 1", 19, 22, 10), ("Heading 2", 14, 15, 6)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Some Word defaults add a colored rule to the Title style. The book uses
    # typography and whitespace only, so explicitly remove any paragraph border.
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    first_h1 = True
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            text = line[2:]
            if first_h1:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.add_run(text)
                first_h1 = False
            else:
                doc.add_heading(text, level=1)
            i += 1
            continue
        if line.startswith("## "):
            text = line[3:]
            if len(doc.paragraphs) == 1:
                doc.add_paragraph(text, style="Subtitle")
            else:
                doc.add_heading(text, level=2)
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= {"-", ":"} for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.autofit = True
            for r_idx, row in enumerate(rows):
                for c_idx, value in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = value
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_margins(cell)
                    if r_idx == 0:
                        set_cell_shading(cell, "315B46")
                        for run in cell.paragraphs[0].runs:
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            run.bold = True
                    elif r_idx % 2 == 0:
                        set_cell_shading(cell, "EEF4F0")
            set_repeat_table_header(table.rows[0])
            doc.add_paragraph()
            continue
        if re.match(r"^\d+\. ", line):
            # Preserve the source number so each independent option set can
            # restart at 1 instead of inheriting Word's document-wide counter.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            add_inline(p, line)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        i += 1

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.text = "Livre Vert de la Confédération des communs terrestres   "
        footer.runs[0].font.size = Pt(8)
        footer.runs[0].font.color.rgb = RGBColor(90, 90, 90)
        add_page_number(footer)

    core = doc.core_properties
    core.title = "Livre Vert de la Confédération des communs terrestres"
    core.subject = "Consultation sur les choix institutionnels et les épreuves à conduire"
    core.author = "Projet de recherche CCT"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == "__main__":
    build()
